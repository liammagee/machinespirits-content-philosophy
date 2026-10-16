from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).resolve().parents[3]
source=(root/'courses/590-fall-2026/syllabus.md').read_text()
source=re.sub(r'^---.*?---\s*','',source,flags=re.S)
parts=re.split(r'^## (.+)\n',source,flags=re.M)
sections={parts[i]:parts[i+1].strip() for i in range(1,len(parts),2)}
intro=sections.pop('ERAM 590 · Fall 2026 Term B · 4 credit hours').replace('  \n','\n\n')
d=Document(root/'courses/479-fall-2026/syllabus/University of Illinois Urbana-Champaign Syllabus 479 Fall 2026.docx')
body=d.element.body
for e in list(body):
 if e.tag!=qn('w:sectPr'):body.remove(e)
for sec in d.sections:
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 for container in [sec.header,sec.footer]:
  for p in container.paragraphs:p.clear()
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run('ERAM 590 · Fall 2026 · Draft  |  ');r.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for name in ['Normal','Body Text','List Paragraph','List Bullet']:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.line_spacing=1.08;st.paragraph_format.widow_control=True
for name,size in [('Title',22),('Subtitle',12),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.font.bold=name!='Subtitle';st.paragraph_format.space_before=Pt(14 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
 st.paragraph_format.keep_with_next=True
 for x in st.element.xpath('.//w:color'):
  for a in list(x.attrib):
   if a!=qn('w:val'):del x.attrib[a]
 for x in st.element.xpath('.//w:pBdr'):x.getparent().remove(x)

def hyperlink(p,label,url):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text):
 pattern=r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|<[^>]+@[^>]+>)'
 for s in re.split(pattern,text):
  if not s:continue
  if re.match(r'^\[[^\]]+\]\(',s):
   label,url=re.match(r'^\[([^\]]+)\]\((.+)\)$',s).groups();hyperlink(p,label,url)
  else:
   bold=s.startswith('**');italic=s.startswith('*') and not bold
   if bold:s=s[2:-2]
   elif italic:s=s[1:-1]
   elif s.startswith('`'):s=s[1:-1]
   elif s.startswith('<') and s.endswith('>'):s=s[1:-1]
   r=p.add_run(s);r.bold=bold;r.italic=italic
   if '[TO CONFIRM]' in s:r.font.highlight_color=7

def table(lines):
 rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
 rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',c) for c in row)]
 n=len(rows[0]);t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={2:[2.0,4.5],3:[3.55,.65,2.3],5:[.75,1.05,1.45,1.8,1.45]}.get(n,[6.5/n]*n)
 if n==2 and rows[0][0]=='Grade':widths=[1.7,2.3]
 if n==2 and rows[0][0]=='Activity':widths=[4.9,1.6]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,(c,text) in enumerate(zip(cells,row)):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'45' if rows[0][0]=='Grade' and side in ['top','bottom'] else '90');x.set(qn('w:type'),'dxa');m.append(x)
   pr.append(m)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECF0');pr.append(shade)
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.0
   inline(p,text)
   for r in p.runs:r.font.size=Pt(9.5 if n==5 else 10);r.bold=True if i==0 else r.bold
  if i==0:
   repeat=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(repeat)
  cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
 d.add_paragraph().paragraph_format.space_after=Pt(1)

def markdown(text):
 lines=text.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   group=[]
   while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
   table(group);continue
  if line.startswith('### '):d.add_heading(line[4:],level=2)
  elif line.startswith('- [ ] '):inline(d.add_paragraph(style='List Bullet'),'[ ] '+line[6:])
  elif line.startswith('- '):inline(d.add_paragraph(style='List Bullet'),line[2:])
  elif re.match(r'^\d+\. ',line):inline(d.add_paragraph(style='List Paragraph'),line)
  else:
   text=line.removeprefix('> ')
   while i+1<len(lines) and lines[i+1].strip() and not re.match(r'^(#|\||- |\d+\. )',lines[i+1].strip()):
    i+=1;text+=' '+lines[i].strip().removeprefix('> ')
   inline(d.add_paragraph(),text)
  i+=1

def section(title,key):d.add_heading(title,level=2);markdown(sections[key])
d.add_paragraph('AI and Academic Literature',style='Title')
d.add_paragraph('ERAM 590 · Fall 2026 Term B · Syllabus Draft',style='Subtitle')
d.add_paragraph('University of Illinois Urbana-Champaign',style='Normal')
d.add_heading('Course Information',level=1)
markdown(intro)
section('Course Description','Course description')
section('Learning Outcomes','Course objectives')
section('Course Outline','Course outline')
d.add_heading('Course Materials',level=1)
markdown(sections['Texts readings and tools'])
d.add_heading('Course Requirements and Policies',level=1)
section('What I Expect from You','What I expect from you')
section('What You Should Expect from Me','What you should expect from me')
section('Evaluation','Evaluation')
section('Grading Scale','Grading scale')
d.add_heading('Campus Academic Policies and Course Arrangements',level=1)
markdown(sections['Policies'])
d.add_heading('Course Schedule and Milestones',level=1)
markdown(sections['Weekly schedule'])
d.add_heading('Instructor Completion Checklist',level=1)
markdown(sections['Instructor completion checklist'])
d.add_heading('Adaptation and Calendar Sources',level=1)
markdown(sections['Adaptation and calendar sources'])
d.core_properties.title='ERAM 590 AI and Academic Literature Syllabus Draft'
d.core_properties.author='Liam Magee';d.core_properties.subject='Fall 2026 Term B'
d.core_properties.comments=''
out=Path(__file__).parent/'University-of-Illinois-Urbana-Champaign-Syllabus-590-Fall-2026-Draft.docx'
d.save(out)
print(out)
